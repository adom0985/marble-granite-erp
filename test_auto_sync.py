import urllib.request, json
import openpyxl
import sqlite3

# 1. Post a test payment of 5,000 EGP for الحاج صلاح
url = 'http://127.0.0.1:8000/api/payments'
payload = {
    'client_name': 'الحاج صلاح',
    'payment_date': '2026-10-27',
    'amount': 5000.0,
    'currency': 'EGP',
    'exchange_rate': 1.0,
    'payment_method': 'نقداً',
    'notes': 'دفعة تجريبية واصل منه'
}
req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers={'Content-Type': 'application/json'})
resp = urllib.request.urlopen(req)
print('Payment API response:', resp.read().decode('utf-8'))

# 2. Check Statement API
url_stmt = 'http://127.0.0.1:8000/api/clients/%D8%A7%D9%84%D8%AD%D8%A7%D8%AC%20%D8%B5%D9%84%D8%A7%D8%AD/statement'
req_stmt = urllib.request.urlopen(url_stmt)
data = json.loads(req_stmt.read().decode('utf-8'))
print('After payment -> Total Paid:', data['total_payments'], 'Remaining:', data['net_remaining'])

# 3. Check Excel file on disk
wb = openpyxl.load_workbook(r'I:\خالي أحمد جلال\الحاج صلاح\حساب الحاج صلاح - النسخة المطورة 2026.xlsx', data_only=False)
ws = wb['الحاج صلاح جديد']
for r in range(11, ws.max_row+1):
    vals = [ws.cell(r, c).value for c in range(1, 6)]
    if any(vals):
        print(f'Excel Row {r}:', vals)

# 4. Rollback the test payment
conn = sqlite3.connect(r'C:\Users\الادهم\.gemini\antigravity\scratch\marble_granite_erp\data\marble_erp.db')
c = conn.cursor()
c.execute("DELETE FROM client_payments WHERE notes = 'دفعة تجريبية واصل منه'")
conn.commit()
conn.close()

from client_statement_sync import sync_client_statement_excel
sync_client_statement_excel('الحاج صلاح')
print('Rolled back test payment cleanly!')
