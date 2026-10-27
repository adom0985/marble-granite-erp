import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def create_invoice_excel(invoice_data, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ورقة1"
    ws.sheet_view.rightToLeft = True

    # Styling definitions
    font_main = Font(name="Arial", size=14, bold=True)
    font_header = Font(name="Arial", size=14, bold=True)
    font_bold_lg = Font(name="Arial", size=16, bold=True)
    font_normal = Font(name="Arial", size=12, bold=True)

    fill_yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    fill_blue_hdr = PatternFill(start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid")
    fill_green_hdr = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")
    fill_total_red = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    fill_free_item = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="FF000000")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Row 1: Header metadata
    ws.cell(row=1, column=2, value="التاريخ").fill = fill_yellow
    ws.cell(row=1, column=3, value=invoice_data.get("invoice_date", ""))
    ws.cell(row=1, column=4, value="العميل").fill = fill_yellow
    ws.cell(row=1, column=5, value=invoice_data.get("client_name", ""))
    ws.cell(row=1, column=7, value="السائق").fill = fill_yellow
    ws.cell(row=1, column=8, value=invoice_data.get("driver_name", ""))

    # Row 2: Secondary metadata
    ws.cell(row=2, column=2, value="رقم السيارة / الكونتر").fill = fill_yellow
    ws.cell(row=2, column=3, value=invoice_data.get("car_no", ""))
    ws.cell(row=2, column=4, value="مكان التحميل").fill = fill_yellow
    ws.cell(row=2, column=5, value=invoice_data.get("loading_place", ""))
    ws.cell(row=2, column=7, value="المشرف").fill = fill_yellow
    ws.cell(row=2, column=8, value=invoice_data.get("supervisor", "-"))
    ws.cell(row=2, column=9, value="رقم الفاتوره").fill = fill_yellow
    ws.cell(row=2, column=10, value=invoice_data.get("invoice_no", ""))

    for r in range(1, 3):
        for c in range(2, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = font_main
            cell.border = cell_border
            cell.alignment = align_center

    # Row 3: Items Grid Header
    headers = [
        "رقم\nالطرد", "رقم البلوك", "نوع الخامه", "عدد", "طول", "عرض", "المسطح", "السعر", "الإجمالي", "ملاحظات"
    ]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_blue_hdr
        cell.border = cell_border
        cell.alignment = align_center

    # Data Rows
    items = invoice_data.get("items", [])
    start_row = 4
    current_row = start_row

    for it in items:
        is_free = bool(it.get("is_free", False))
        ws.cell(row=current_row, column=1, value=it.get("package_no", ""))
        ws.cell(row=current_row, column=2, value=it.get("block_no", ""))
        ws.cell(row=current_row, column=3, value=it.get("material_name", ""))
        ws.cell(row=current_row, column=4, value=float(it.get("qty", 0)))
        ws.cell(row=current_row, column=5, value=float(it.get("length", 0)))
        ws.cell(row=current_row, column=6, value=float(it.get("width", 0)))
        
        # Area = PRODUCT(D:F)
        ws.cell(row=current_row, column=7, value=f"=PRODUCT(D{current_row}:F{current_row})")
        
        if is_free:
            ws.cell(row=current_row, column=8, value=0)
            ws.cell(row=current_row, column=9, value=0)
            note_val = str(it.get("notes", "") or "")
            if "فري" not in note_val:
                note_val = ("فري (مجاني) " + note_val).strip()
            ws.cell(row=current_row, column=10, value=note_val)
        else:
            ws.cell(row=current_row, column=8, value=float(it.get("price", 0)))
            ws.cell(row=current_row, column=9, value=f"=PRODUCT(G{current_row}:H{current_row})")
            ws.cell(row=current_row, column=10, value=it.get("notes", ""))

        for c in range(1, 11):
            cell = ws.cell(row=current_row, column=c)
            cell.font = font_normal
            cell.border = cell_border
            cell.alignment = align_center
            if is_free:
                cell.fill = fill_free_item
            if c in [4, 5, 6, 7, 8, 9]:
                cell.number_format = "#,##0.00"

        current_row += 1

    end_row = current_row - 1 if current_row > start_row else start_row

    # Total Row
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    cell_tot_lbl = ws.cell(row=current_row, column=1, value="الإجمالي")
    cell_tot_lbl.font = font_bold_lg
    cell_tot_lbl.fill = fill_total_red
    cell_tot_lbl.alignment = align_center

    cell_tot_qty = ws.cell(row=current_row, column=4, value=f"=SUM(D{start_row}:D{end_row})")
    cell_tot_qty.font = font_bold_lg
    cell_tot_qty.fill = fill_total_red
    cell_tot_qty.alignment = align_center

    for c in [5, 6, 8, 10]:
        ws.cell(row=current_row, column=c, value="").fill = fill_total_red

    cell_tot_area = ws.cell(row=current_row, column=7, value=f"=SUM(G{start_row}:G{end_row})")
    cell_tot_area.font = font_bold_lg
    cell_tot_area.fill = fill_total_red
    cell_tot_area.alignment = align_center
    cell_tot_area.number_format = "#,##0.00"

    cell_tot_val = ws.cell(row=current_row, column=9, value=f"=SUM(I{start_row}:I{end_row})")
    cell_tot_val.font = font_bold_lg
    cell_tot_val.fill = fill_total_red
    cell_tot_val.alignment = align_center
    cell_tot_val.number_format = "#,##0.00"

    for c in range(1, 11):
        ws.cell(row=current_row, column=c).border = cell_border

    # الملخص (ملخص الجرار)
    summary_start = current_row + 3
    ws.cell(row=summary_start, column=3, value="ملخص الجرار").fill = fill_yellow
    ws.cell(row=summary_start, column=3).font = font_bold_lg
    ws.cell(row=summary_start, column=3).alignment = align_center
    ws.cell(row=summary_start, column=3).border = cell_border

    sum_hdr_row = summary_start + 1
    sum_headers = [("الخامه و البيان", 3), ("عدد الطاولات", 4), ("كمية الأمتار", 5), ("السعر", 6), ("الإجمالي", 7)]
    for title, col_idx in sum_headers:
        cell = ws.cell(row=sum_hdr_row, column=col_idx, value=title)
        cell.fill = fill_green_hdr
        cell.font = font_header
        cell.alignment = align_center
        cell.border = cell_border

    # Aggregate items strictly by material (no length or width split)
    agg = {}
    for it in items:
        mat = (it.get("material_name") or "خامة").strip()
        is_free = bool(it.get("is_free", False))
        price = 0 if is_free else float(it.get("price") or 0)
        suffix = " ( فري )" if is_free else ""
        key = f"{mat}{suffix}"
        if key not in agg:
            agg[key] = {"qty": 0.0, "area": 0.0, "price": price}
        agg[key]["qty"] += float(it.get("qty") or 0)
        agg[key]["area"] += float(it.get("area") or 0)
        if not is_free and price > 0:
            agg[key]["price"] = price

    s_row = sum_hdr_row + 1
    for key, val in agg.items():
        ws.cell(row=s_row, column=3, value=key).alignment = align_center
        ws.cell(row=s_row, column=4, value=val["qty"]).alignment = align_center
        ws.cell(row=s_row, column=5, value=round(val["area"], 3)).alignment = align_center
        ws.cell(row=s_row, column=6, value=val["price"]).alignment = align_center
        ws.cell(row=s_row, column=7, value=f"=PRODUCT(E{s_row}:F{s_row})").alignment = align_center
        
        for c in range(3, 8):
            cell = ws.cell(row=s_row, column=c)
            cell.font = font_normal
            cell.border = cell_border
            if c in [4, 5, 6, 7]:
                cell.number_format = "#,##0.00"
        s_row += 1

    # Summary Total Row
    sum_end = s_row - 1 if s_row > sum_hdr_row + 1 else sum_hdr_row + 1
    ws.cell(row=s_row, column=3, value="الإجمالي").fill = fill_green_hdr
    ws.cell(row=s_row, column=4, value=f"=SUM(D{sum_hdr_row+1}:D{sum_end})").fill = fill_green_hdr
    ws.cell(row=s_row, column=5, value=f"=SUM(E{sum_hdr_row+1}:E{sum_end})").fill = fill_green_hdr
    ws.cell(row=s_row, column=6, value="").fill = fill_green_hdr
    ws.cell(row=s_row, column=7, value=f"=SUM(G{sum_hdr_row+1}:G{sum_end})").fill = fill_green_hdr

    for c in range(3, 8):
        cell = ws.cell(row=s_row, column=c)
        cell.font = font_bold_lg
        cell.border = cell_border
        cell.alignment = align_center
        if c in [4, 5, 7]:
            cell.number_format = "#,##0.00"

    col_widths = {
        1: 14, 2: 14, 3: 28, 4: 12, 5: 12, 6: 12, 7: 15, 8: 12, 9: 18, 10: 20
    }
    for c_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
