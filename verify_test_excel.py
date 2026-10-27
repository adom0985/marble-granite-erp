import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\الادهم\.gemini\antigravity\scratch\marble_granite_erp\exports\test_invoice.xlsx")
ws = wb.active
print(f"Title: {ws.title}, RTL: {ws.sheet_view.rightToLeft}")
for r in range(1, ws.max_row + 1):
    vals = [f"C{c}:{ws.cell(row=r, column=c).value}" for c in range(1, 11) if ws.cell(row=r, column=c).value is not None]
    if vals:
        print(f"R{r:2d}: " + " | ".join(vals))
