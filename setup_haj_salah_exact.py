import sqlite3
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil

DB_PATH = r"C:\Users\الادهم\.gemini\antigravity\scratch\marble_granite_erp\data\marble_erp.db"

# 1. Update SQLite Database
conn = sqlite3.connect(DB_PATH)
c = conn.cursor()

# Remove test invoice 2
c.execute("DELETE FROM invoices WHERE id = 2")
c.execute("DELETE FROM invoice_items WHERE invoice_id = 2")

# Update active invoices
active_updates = [
    (13, "فاتورة 23 / 9 (1)", "23/9/2026", 121755.0, "حساب الحاج صلاح جديد"),
    (12, "فاتورة 23 / 9 (2)", "23/9/2026", 122755.0, "حساب الحاج صلاح جديد"),
    (15, "فاتورة 3 / 10", "3/10/2026", 125242.0, "حساب الحاج صلاح جديد"),
    (3,  "فاتورة 11 / 10", "11/10/2026", 122828.0, "حساب الحاج صلاح جديد")
]

for inv_id, inv_no, inv_date, amount, note in active_updates:
    c.execute("""
    UPDATE invoices 
    SET invoice_no = ?, invoice_date = ?, total_amount = ?, net_amount = ?, 
        paid_amount = 0.0, remaining_amount = ?, status = 'جارية', notes = ?
    WHERE id = ?
    """, (inv_no, inv_date, amount, amount, amount, note, inv_id))

# Update all other historical invoices for الحاج صلاح to closed archive
c.execute("""
UPDATE invoices 
SET status = 'متقفلة (أرشيف قديم)', paid_amount = total_amount, remaining_amount = 0.0, notes = 'فاتورة سابقة متقفلة ومسددة'
WHERE client_name = 'الحاج صلاح' AND id NOT IN (13, 12, 15, 3)
""")

# Clean previous test payments for الحاج صلاح
c.execute("DELETE FROM client_payments WHERE client_name = 'الحاج صلاح'")

# Insert the two exact payments from the user's sheet
c.execute("""
INSERT INTO client_payments (client_name, invoice_id, payment_date, amount, currency, exchange_rate, amount_egp, payment_method, notes)
VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
""", ("الحاج صلاح", None, "2026-10-15", 150000.0, "EGP", 1.0, 150000.0, "نقداً", "واصل منه"))

c.execute("""
INSERT INTO client_payments (client_name, invoice_id, payment_date, amount, currency, exchange_rate, amount_egp, payment_method, notes)
VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
""", ("الحاج صلاح", None, "2026-10-20", 132500.0, "EGP", 1.0, 132500.0, "نقداً", "واصل منه"))

conn.commit()

# Verify calculations
c.execute("SELECT SUM(net_amount) FROM invoices WHERE client_name = 'الحاج صلاح' AND status = 'جارية'")
tot_active = c.fetchone()[0] or 0.0

c.execute("SELECT SUM(amount_egp) FROM client_payments WHERE client_name = 'الحاج صلاح'")
tot_paid = c.fetchone()[0] or 0.0

rem = tot_active - tot_paid
print(f"DB Updated: Active Invoices: {tot_active:,.2f} EGP | Paid: {tot_paid:,.2f} EGP | Remaining: {rem:,.2f} EGP")
conn.close()

# 2. Build Upgraded Excel Workbook for الحاج صلاح
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Font & Styles
font_title = Font(name="Cairo", size=18, bold=True, color="1E3A8A")
font_sub = Font(name="Cairo", size=11, bold=True, color="475569")
font_head = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Cairo", size=11, bold=False, color="0F172A")
font_bold = Font(name="Cairo", size=12, bold=True, color="0F172A")
font_rem = Font(name="Cairo", size=14, bold=True, color="B91C1C")

fill_head = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_green_head = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
fill_tot = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
fill_rem = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
double_bottom = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='double', color='1E3A8A')
)

align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# ----------------- SHEET 1: الحاج صلاح جديد -----------------
ws1 = wb.create_sheet(title="الحاج صلاح جديد")
ws1.views.sheetView[0].rightToLeft = True

# Title
ws1.merge_cells("A1:E1")
ws1["A1"] = "الحاج صلاح جديد"
ws1["A1"].font = font_title
ws1["A1"].alignment = align_center

ws1.merge_cells("A2:E2")
ws1["A2"] = "كشف الحساب الجاري والفواتير النشطة والدفعات المسددة"
ws1["A2"].font = font_sub
ws1["A2"].alignment = align_center

# Invoices Header
headers_inv = ["م", "بيان الفاتورة", "المبلغ (ج.م)", "التاريخ", "ملاحظات"]
for col_idx, h in enumerate(headers_inv, 1):
    cell = ws1.cell(row=4, column=col_idx, value=h)
    cell.font = font_head
    cell.fill = fill_head
    cell.alignment = align_center
    cell.border = thin_border

inv_rows = [
    (1, "فاتورة 23 / 9", 121755, "23/9/2026", "طاولات وجرارات"),
    (2, "فاتورة 23 / 9", 122755, "23/9/2026", "رخام وترابيع"),
    (3, "فاتورة 3 / 10", 125242, "3/10/2026", "طاولات صني"),
    (4, "فاتورة 11 / 10", 122828, "11/10/2026", "توريد 3 سم")
]

start_inv_r = 5
for idx, (m, name, amt, dt, note) in enumerate(inv_rows):
    curr_r = start_inv_r + idx
    ws1.cell(row=curr_r, column=1, value=m).alignment = align_center
    ws1.cell(row=curr_r, column=2, value=name).alignment = align_right
    c_amt = ws1.cell(row=curr_r, column=3, value=amt)
    c_amt.alignment = align_center
    c_amt.number_format = '#,##0'
    ws1.cell(row=curr_r, column=4, value=dt).alignment = align_center
    ws1.cell(row=curr_r, column=5, value=note).alignment = align_right
    for c in range(1, 6):
        ws1.cell(row=curr_r, column=c).font = font_data
        ws1.cell(row=curr_r, column=c).border = thin_border

end_inv_r = start_inv_r + len(inv_rows) - 1
tot_inv_r = end_inv_r + 1

# Total Invoices Row
ws1.merge_cells(start_row=tot_inv_r, start_column=1, end_row=tot_inv_r, end_column=2)
ws1.cell(row=tot_inv_r, column=1, value="الاجمالي").alignment = align_center
ws1.cell(row=tot_inv_r, column=1).font = font_bold
ws1.cell(row=tot_inv_r, column=1).fill = fill_tot
ws1.cell(row=tot_inv_r, column=2).fill = fill_tot

c_tot = ws1.cell(row=tot_inv_r, column=3, value=f"=SUM(C{start_inv_r}:C{end_inv_r})")
c_tot.font = font_bold
c_tot.fill = fill_tot
c_tot.alignment = align_center
c_tot.number_format = '#,##0'

for c in range(1, 6):
    ws1.cell(row=tot_inv_r, column=c).border = double_bottom
    if c > 3:
        ws1.cell(row=tot_inv_r, column=c).fill = fill_tot

# Payments Section
pay_head_r = tot_inv_r + 2
headers_pay = ["م", "بيان الدفعة الواصلة", "المبلغ (ج.م)", "التاريخ", "طريقة الدفع"]
for col_idx, h in enumerate(headers_pay, 1):
    cell = ws1.cell(row=pay_head_r, column=col_idx, value=h)
    cell.font = font_head
    cell.fill = fill_green_head
    cell.alignment = align_center
    cell.border = thin_border

pay_rows = [
    (1, "واصل منه", 150000, "15/10/2026", "نقداً"),
    (2, "واصل منه", 132500, "20/10/2026", "نقداً")
]

start_pay_r = pay_head_r + 1
for idx, (m, name, amt, dt, meth) in enumerate(pay_rows):
    curr_r = start_pay_r + idx
    ws1.cell(row=curr_r, column=1, value=m).alignment = align_center
    ws1.cell(row=curr_r, column=2, value=name).alignment = align_right
    c_amt = ws1.cell(row=curr_r, column=3, value=amt)
    c_amt.alignment = align_center
    c_amt.number_format = '#,##0'
    ws1.cell(row=curr_r, column=4, value=dt).alignment = align_center
    ws1.cell(row=curr_r, column=5, value=meth).alignment = align_center
    for c in range(1, 6):
        ws1.cell(row=curr_r, column=c).font = font_data
        ws1.cell(row=curr_r, column=c).border = thin_border

end_pay_r = start_pay_r + len(pay_rows) - 1
tot_pay_r = end_pay_r + 1

# Total Payments Row
ws1.merge_cells(start_row=tot_pay_r, start_column=1, end_row=tot_pay_r, end_column=2)
ws1.cell(row=tot_pay_r, column=1, value="إجمالي الواصل منه").alignment = align_center
ws1.cell(row=tot_pay_r, column=1).font = font_bold
ws1.cell(row=tot_pay_r, column=1).fill = fill_tot
ws1.cell(row=tot_pay_r, column=2).fill = fill_tot

c_pay_tot = ws1.cell(row=tot_pay_r, column=3, value=f"=SUM(C{start_pay_r}:C{end_pay_r})")
c_pay_tot.font = font_bold
c_pay_tot.fill = fill_tot
c_pay_tot.alignment = align_center
c_pay_tot.number_format = '#,##0'

for c in range(1, 6):
    ws1.cell(row=tot_pay_r, column=c).border = double_bottom
    if c > 3:
        ws1.cell(row=tot_pay_r, column=c).fill = fill_tot

# Remaining Balance Row (المتبقي)
rem_r = tot_pay_r + 2
ws1.merge_cells(start_row=rem_r, start_column=1, end_row=rem_r, end_column=2)
ws1.cell(row=rem_r, column=1, value="المتبقي (الصافي المطلوب)").alignment = align_center
ws1.cell(row=rem_r, column=1).font = font_rem
ws1.cell(row=rem_r, column=1).fill = fill_rem
ws1.cell(row=rem_r, column=2).fill = fill_rem

c_rem = ws1.cell(row=rem_r, column=3, value=f"=C{tot_inv_r}-C{tot_pay_r}")
c_rem.font = font_rem
c_rem.fill = fill_rem
c_rem.alignment = align_center
c_rem.number_format = '#,##0'

for c in range(1, 6):
    ws1.cell(row=rem_r, column=c).border = double_bottom
    if c > 3:
        ws1.cell(row=rem_r, column=c).fill = fill_rem

# Set Column Widths
ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['B'].width = 28
ws1.column_dimensions['C'].width = 20
ws1.column_dimensions['D'].width = 16
ws1.column_dimensions['E'].width = 25

# ----------------- SHEET 2: جديد (Exact Replica of User Sheet) -----------------
ws2 = wb.create_sheet(title="جديد")
ws2.views.sheetView[0].rightToLeft = True

ws2["B1"] = "الحاج صلاح جديد"
ws2["B1"].font = font_title
ws2["B1"].alignment = align_center

ws2["A4"] = "فاتورة 23 / 9 "
ws2["B4"] = 121755
ws2["C4"] = "23/9/2026"

ws2["A5"] = "فاتورة 23 / 9 "
ws2["B5"] = 122755
ws2["C5"] = "23/9/2026"

ws2["A6"] = "فاتورة 3 / 10"
ws2["B6"] = 125242
ws2["C6"] = "3/10/2026"

ws2["A7"] = "فاتورة 11 / 10"
ws2["B7"] = 122828
ws2["C7"] = "11/10/2026"

ws2["A9"] = "الاجمالي"
ws2["B9"] = "=SUM(B4:B7)"

ws2["A11"] = "واصل منه"
ws2["B11"] = 150000

ws2["A12"] = "واصل منه"
ws2["B12"] = 132500

ws2["A14"] = "المتبقى"
ws2["B14"] = "=SUM(B9-B11-B12)"

for r in [4, 5, 6, 7, 9, 11, 12, 14]:
    ws2[f"A{r}"].font = font_bold
    ws2[f"B{r}"].font = font_bold
    ws2[f"B{r}"].number_format = '#,##0'
    ws2[f"B{r}"].alignment = align_center
    if ws2[f"C{r}"].value:
        ws2[f"C{r}"].font = font_data
        ws2[f"C{r}"].alignment = align_center

ws2["A9"].fill = fill_tot
ws2["B9"].fill = fill_tot
ws2["A14"].fill = fill_rem
ws2["B14"].fill = fill_rem

ws2.column_dimensions['A'].width = 22
ws2.column_dimensions['B'].width = 20
ws2.column_dimensions['C'].width = 18

# ----------------- SHEET 3: الأرشيف التاريخي المقفل -----------------
ws3 = wb.create_sheet(title="الأرشيف القديم (متقفل)")
ws3.views.sheetView[0].rightToLeft = True

ws3.merge_cells("A1:F1")
ws3["A1"] = "أرشيف الفواتير السابقة المقفلة - الحاج صلاح"
ws3["A1"].font = font_title
ws3["A1"].alignment = align_center

headers_arch = ["م", "رقم الفاتورة", "التاريخ", "المبلغ (ج.م)", "الحالة", "ملاحظات"]
for col_idx, h in enumerate(headers_arch, 1):
    cell = ws3.cell(row=3, column=col_idx, value=h)
    cell.font = font_head
    cell.fill = fill_head
    cell.alignment = align_center
    cell.border = thin_border

conn = sqlite3.connect(DB_PATH)
c = conn.cursor()
c.execute("SELECT id, invoice_no, invoice_date, net_amount, status, notes FROM invoices WHERE client_name = 'الحاج صلاح' AND status = 'متقفلة (أرشيف قديم)' ORDER BY id ASC")
arch_invoices = c.fetchall()
conn.close()

for idx, (inv_id, inv_no, inv_date, net_amt, st, note) in enumerate(arch_invoices, 1):
    r = 3 + idx
    ws3.cell(row=r, column=1, value=idx).alignment = align_center
    ws3.cell(row=r, column=2, value=inv_no).alignment = align_right
    ws3.cell(row=r, column=3, value=str(inv_date)).alignment = align_center
    c_amt = ws3.cell(row=r, column=4, value=net_amt)
    c_amt.alignment = align_center
    c_amt.number_format = '#,##0.00'
    ws3.cell(row=r, column=5, value="خالص ومقفل").alignment = align_center
    ws3.cell(row=r, column=6, value=note or "").alignment = align_right
    for col in range(1, 7):
        ws3.cell(row=r, column=col).font = font_data
        ws3.cell(row=r, column=col).border = thin_border

ws3.column_dimensions['A'].width = 8
ws3.column_dimensions['B'].width = 30
ws3.column_dimensions['C'].width = 16
ws3.column_dimensions['D'].width = 20
ws3.column_dimensions['E'].width = 16
ws3.column_dimensions['F'].width = 35

# Save files:
path_upgraded = r"I:\خالي أحمد جلال\الحاج صلاح\حساب الحاج صلاح - النسخة المطورة 2026.xlsx"
path_jadeed = r"I:\خالي أحمد جلال\الحاج صلاح\حساب الحاج صلاح جديد 2.xlsx"
path_bundled = r"C:\Users\الادهم\.gemini\antigravity\scratch\marble_granite_erp\bundled_files\الحاج صلاح\حساب الحاج صلاح - النسخة المطورة 2026.xlsx"

wb.save(path_upgraded)
print(f"Saved: {path_upgraded}")

wb.save(path_bundled)
print(f"Saved: {path_bundled}")

# Also update sheet جديد in path_jadeed
try:
    wb_j = openpyxl.load_workbook(path_jadeed)
    if "جديد" in wb_j.sheetnames:
        del wb_j["جديد"]
    ws_new = wb_j.create_sheet("جديد", 1) # Insert after Sheet1
    ws_new.views.sheetView[0].rightToLeft = True
    for r in range(1, 16):
        for c in range(1, 4):
            val = ws2.cell(r, c).value
            if val is not None:
                ws_new.cell(r, c, val)
    wb_j.save(path_jadeed)
    print(f"Updated 'جديد' inside {path_jadeed}")
except Exception as e:
    print(f"Notice on path_jadeed: {e}")

print("All Excel files and Database aligned with the exact screenshot!")
