import sqlite3
import os, shutil, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

DB_PATH = "data/marble_erp.db"
INVOICE_ID = 87  # فاتورة 27 / 10

# New correct values
TOTAL_BOXES = 28
TILES_PER_BOX = 220
TOTAL_TILES = TOTAL_BOXES * TILES_PER_BOX  # 6160
TILE_LENGTH = 0.60
TILE_WIDTH = 0.30
AREA_PER_TILE = TILE_LENGTH * TILE_WIDTH  # 0.18
TOTAL_AREA = TOTAL_TILES * AREA_PER_TILE  # 1108.80
PRICE_PER_M2 = 185.0
TOTAL_AMOUNT = TOTAL_AREA * PRICE_PER_M2  # 205128.00

print(f"=== Corrected Invoice Calculation ===")
print(f"Boxes: {TOTAL_BOXES}")
print(f"Tiles per box: {TILES_PER_BOX}")
print(f"Total tiles: {TOTAL_TILES}")
print(f"Tile size: {TILE_LENGTH} x {TILE_WIDTH} = {AREA_PER_TILE} m²")
print(f"Total area: {TOTAL_AREA} m²")
print(f"Price/m²: {PRICE_PER_M2}")
print(f"Total: {TOTAL_AMOUNT:,.2f} EGP")
print()

# Update DB
conn = sqlite3.connect(DB_PATH)
conn.row_factory = sqlite3.Row
c = conn.cursor()

# Update invoice totals
c.execute("""UPDATE invoices SET 
    total_pieces = ?,
    total_area = ?,
    total_amount = ?,
    net_amount = ?,
    remaining_amount = ?
    WHERE id = ?""", 
    (TOTAL_TILES, TOTAL_AREA, TOTAL_AMOUNT, TOTAL_AMOUNT, TOTAL_AMOUNT, INVOICE_ID))

# Update invoice item
c.execute("""UPDATE invoice_items SET
    qty = ?,
    length = ?,
    width = ?,
    area = ?,
    price = ?,
    total = ?,
    notes = ?
    WHERE invoice_id = ?""",
    (TOTAL_TILES, TILE_LENGTH, TILE_WIDTH, TOTAL_AREA, PRICE_PER_M2, TOTAL_AMOUNT, 
     f"28 صندوق × 220 تربيعة", INVOICE_ID))

conn.commit()

# Verify
c.execute("SELECT * FROM invoices WHERE id = ?", (INVOICE_ID,))
inv = dict(c.fetchone())
print(f"Updated Invoice {INVOICE_ID}: pieces={inv['total_pieces']}, area={inv['total_area']}, amount={inv['total_amount']}, net={inv['net_amount']}")

c.execute("SELECT * FROM invoice_items WHERE invoice_id = ?", (INVOICE_ID,))
items = [dict(r) for r in c.fetchall()]
for it in items:
    print(f"  Item: qty={it['qty']}, area={it['area']}, price={it['price']}, total={it['total']}, notes={it['notes']}")

# Recalculate active statement for Haj Salah
c.execute("SELECT id, invoice_no, total_amount, net_amount, status FROM invoices WHERE client_name = ? AND status != ?", 
          ("الحاج صلاح", "متقفلة (أرشيف قديم)"))
active_invs = [dict(r) for r in c.fetchall()]
active_total = sum(i['net_amount'] for i in active_invs)

c.execute("SELECT SUM(amount_egp) as total_paid FROM client_payments WHERE client_name = ?", ("الحاج صلاح",))
paid_row = c.fetchone()
total_paid = paid_row['total_paid'] or 0

remaining = active_total - total_paid
print(f"\n=== Updated Haj Salah Statement ===")
print(f"Active invoices: {len(active_invs)}")
for ai in active_invs:
    print(f"  {ai['invoice_no']}: {ai['net_amount']:,.2f}")
print(f"Active total: {active_total:,.2f}")
print(f"Paid: {total_paid:,.2f}")
print(f"Remaining: {remaining:,.2f}")

conn.close()

# Update Excel statement
WORK_DIR = r"I:\خالي أحمد جلال"
CLIENT_DIR = os.path.join(WORK_DIR, "الحاج صلاح")
STATEMENT_FILE = os.path.join(CLIENT_DIR, "حساب الحاج صلاح - النسخة المطورة 2026.xlsx")
BUNDLED_DIR = os.path.join("bundled_files", "الحاج صلاح")

if os.path.exists(STATEMENT_FILE):
    wb = openpyxl.load_workbook(STATEMENT_FILE)
    ws = wb.active
    
    # Row 9 = invoice فاتورة 27 / 10 - update the amount
    ws.cell(row=9, column=3, value=TOTAL_AMOUNT)
    ws.cell(row=9, column=5, value=f"ترابيع سيلفيا {TOTAL_BOXES} صندوق × {TILES_PER_BOX} تربيعة")
    
    wb.save(STATEMENT_FILE)
    print(f"\nUpdated Excel: {STATEMENT_FILE}")
    
    # Sync to bundled
    os.makedirs(BUNDLED_DIR, exist_ok=True)
    bundled_path = os.path.join(BUNDLED_DIR, "حساب الحاج صلاح - النسخة المطورة 2026.xlsx")
    shutil.copy2(STATEMENT_FILE, bundled_path)
    print(f"Synced to: {bundled_path}")
    
    # Verify
    wb2 = openpyxl.load_workbook(STATEMENT_FILE)
    ws2 = wb2.active
    for r in range(5, 18):
        row_data = [ws2.cell(row=r, column=c).value for c in range(1, 6)]
        if any(v is not None for v in row_data):
            print(f"  Row {r}: {row_data}")
else:
    print(f"\nWarning: Statement file not found at {STATEMENT_FILE}")

# Also regenerate the individual invoice Excel+PDF
from excel_exporter import create_invoice_excel
from pdf_invoice_generator import generate_invoice_pdf

conn2 = sqlite3.connect(DB_PATH)
conn2.row_factory = sqlite3.Row
c2 = conn2.cursor()
c2.execute("SELECT * FROM invoices WHERE id = ?", (INVOICE_ID,))
inv_dict = dict(c2.fetchone())
c2.execute("SELECT * FROM invoice_items WHERE invoice_id = ? ORDER BY row_index", (INVOICE_ID,))
inv_dict["items"] = [dict(r) for r in c2.fetchall()]
conn2.close()

xlsx_path = os.path.join(CLIENT_DIR, "2026-10-27 - فاتورة 27-10.xlsx")
pdf_path = os.path.join(CLIENT_DIR, "2026-10-27 - فاتورة 27-10.pdf")

create_invoice_excel(inv_dict, xlsx_path)
generate_invoice_pdf(inv_dict, pdf_path)
print(f"Regenerated: {xlsx_path} ({os.path.getsize(xlsx_path)} bytes)")
print(f"Regenerated: {pdf_path} ({os.path.getsize(pdf_path)} bytes)")

print("\n✅ Invoice corrected successfully!")
