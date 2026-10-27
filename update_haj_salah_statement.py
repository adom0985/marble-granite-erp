import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

HAJ_SALAH_EXCEL = r"I:\خالي أحمد جلال\الحاج صلاح\حساب الحاج صلاح - النسخة المطورة 2026.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "كشف الحساب المطور 2026"
ws.views.sheetView[0].rightToLeft = True

# Styles
font_title = Font(name="Arial", size=16, bold=True, color="1E3A8A")
font_head = Font(name="Arial", size=12, bold=True, color="FFFFFF")
font_bold = Font(name="Arial", size=12, bold=True, color="000000")
font_normal = Font(name="Arial", size=11, bold=False, color="000000")
font_tot = Font(name="Arial", size=13, bold=True, color="000000")
font_rem = Font(name="Arial", size=15, bold=True, color="B91C1C")

fill_head_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_head_green = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
fill_yellow = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
fill_pink = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

thin = Side(border_style="thin", color="000000")
cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
align_c = Alignment(horizontal="center", vertical="center")

# Title Row - NO "شركة"
ws.merge_cells("A1:E1")
ws["A1"] = "كشف حساب وسجل حركات: الحاج صلاح (النسخة المطورة 2026)"
ws["A1"].font = font_title
ws["A1"].alignment = align_c

# Invoices Section
headers_inv = ["م", "بيان الفاتورة", "المبلغ (ج.م)", "التاريخ", "ملاحظات وتفاصيل التوريد"]
for col_idx, h in enumerate(headers_inv, 1):
    c = ws.cell(row=3, column=col_idx, value=h)
    c.font = font_head
    c.fill = fill_head_blue
    c.alignment = align_c
    c.border = cell_border

invoices_data = [
    (1, "فاتورة 11 / 10", 122828.00, "11/10/2026", "137 طاولة صني 3 سم (744.41 م٢)"),
    (2, "فاتورة 23 / 9 (2)", 122755.00, "23/9/2026", "138 طاولة صني 3 سم (743.90 م٢)"),
    (3, "فاتورة 23 / 9 (1)", 121755.00, "23/9/2026", "133 طاولة صني 3 سم (737.90 م٢)"),
    (4, "فاتورة 3 / 10", 125242.00, "3/10/2026", "141 طاولة سيلفيا 3 سم (659.17 م٢)"),
    (5, "فاتورة 27 / 10 (1)", 205128.00, "27/10/2026", "28 صندوق ترابيع سيلفيا = 6,160 تربيعة (1,108.80 م٢)"),
    (6, "فاتورة 26 / 10 (1)", 122985.05, "26/10/2026", "132 طاولة صني 3 سم (745.36 م٢)"),
    (7, "فاتورة 26 / 10 (2)", 122488.81, "26/10/2026", "132 طاولة صني 3 سم (742.36 م٢)"),
    (8, "فاتورة 27 / 10 (2)", 325114.87, "27/10/2026", "360 طاولة جرانيت وسيلفيا (كل الأطوال +7سم و 3 شرائح رويال فري) (919.54 م٢)"),
]

row_idx = 4
for item in invoices_data:
    ws.cell(row=row_idx, column=1, value=item[0])
    ws.cell(row=row_idx, column=2, value=item[1])
    ws.cell(row=row_idx, column=3, value=item[2]).number_format = "#,##0.00"
    ws.cell(row=row_idx, column=4, value=item[3])
    ws.cell(row=row_idx, column=5, value=item[4])
    
    for col in range(1, 6):
        c = ws.cell(row=row_idx, column=col)
        c.font = font_normal if col != 3 else font_bold
        c.alignment = align_c
        c.border = cell_border
    row_idx += 1

# Total Invoices
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
ws.cell(row=row_idx, column=1, value="إجمالي المسحوبات (الفواتير)").font = font_tot
ws.cell(row=row_idx, column=1).alignment = align_c
ws.cell(row=row_idx, column=1).fill = fill_yellow

ws.cell(row=row_idx, column=3, value=f"=SUM(C4:C{row_idx-1})").font = font_tot
ws.cell(row=row_idx, column=3).alignment = align_c
ws.cell(row=row_idx, column=3).fill = fill_yellow
ws.cell(row=row_idx, column=3).number_format = "#,##0.00"

for col in range(1, 6):
    ws.cell(row=row_idx, column=col).border = cell_border
    ws.cell(row=row_idx, column=col).fill = fill_yellow
row_tot_inv = row_idx

# Payments Section
row_idx += 2
headers_pay = ["م", "بيان الدفعة الواصلة", "المبلغ (ج.م)", "التاريخ", "طريقة الدفع"]
for col_idx, h in enumerate(headers_pay, 1):
    c = ws.cell(row=row_idx, column=col_idx, value=h)
    c.font = font_head
    c.fill = fill_head_green
    c.alignment = align_c
    c.border = cell_border

pay_start = row_idx + 1
payments_data = [
    (1, "واصل منه", 150000.00, "15/10/2026", "نقداً"),
    (2, "واصل منه", 132500.00, "20/10/2026", "نقداً"),
]

row_idx += 1
for p in payments_data:
    ws.cell(row=row_idx, column=1, value=p[0])
    ws.cell(row=row_idx, column=2, value=p[1])
    ws.cell(row=row_idx, column=3, value=p[2]).number_format = "#,##0.00"
    ws.cell(row=row_idx, column=4, value=p[3])
    ws.cell(row=row_idx, column=5, value=p[4])
    
    for col in range(1, 6):
        c = ws.cell(row=row_idx, column=col)
        c.font = font_normal if col != 3 else font_bold
        c.alignment = align_c
        c.border = cell_border
    row_idx += 1

# Total Payments
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
ws.cell(row=row_idx, column=1, value="إجمالي الواصل منه (السدادات)").font = font_tot
ws.cell(row=row_idx, column=1).alignment = align_c
ws.cell(row=row_idx, column=1).fill = fill_yellow

ws.cell(row=row_idx, column=3, value=f"=SUM(C{pay_start}:C{row_idx-1})").font = font_tot
ws.cell(row=row_idx, column=3).alignment = align_c
ws.cell(row=row_idx, column=3).fill = fill_yellow
ws.cell(row=row_idx, column=3).number_format = "#,##0.00"

for col in range(1, 6):
    ws.cell(row=row_idx, column=col).border = cell_border
    ws.cell(row=row_idx, column=col).fill = fill_yellow
row_tot_pay = row_idx

# Remaining Balance Section
row_idx += 2
ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
ws.cell(row=row_idx, column=1, value="المتبقي (الصافي المطلوب)").font = font_rem
ws.cell(row=row_idx, column=1).alignment = align_c
ws.cell(row=row_idx, column=1).fill = fill_pink

ws.cell(row=row_idx, column=3, value=f"=C{row_tot_inv}-C{row_tot_pay}").font = font_rem
ws.cell(row=row_idx, column=3).alignment = align_c
ws.cell(row=row_idx, column=3).fill = fill_pink
ws.cell(row=row_idx, column=3).number_format = "#,##0.00"

for col in range(1, 6):
    ws.cell(row=row_idx, column=col).border = cell_border
    ws.cell(row=row_idx, column=col).fill = fill_pink

# Column widths
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 20
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 65

wb.save(HAJ_SALAH_EXCEL)
print(f"Master statement saved cleanly without any 'شركة': {HAJ_SALAH_EXCEL}")
