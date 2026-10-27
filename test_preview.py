import os
import openpyxl

target_dir = r"I:\خالي أحمد جلال"

def preview_excel(rel_path):
    full_path = os.path.join(target_dir, rel_path)
    if not os.path.exists(full_path):
        return {"error": "File not found"}
    try:
        wb = openpyxl.load_workbook(full_path, data_only=True)
        sheets_data = {}
        for s in wb.sheetnames[:3]:
            ws = wb[s]
            rows = []
            for r in range(1, min(ws.max_row + 1, 35)):
                row_vals = [str(ws.cell(row=r, column=c).value or "") for c in range(1, min(ws.max_column + 1, 15))]
                if any(v.strip() for v in row_vals):
                    rows.append(row_vals)
            sheets_data[s] = rows
        return {
            "file": os.path.basename(rel_path),
            "sheets": wb.sheetnames,
            "sheets_data": sheets_data
        }
    except Exception as e:
        return {"error": str(e)}

res = preview_excel(r"أحمد جلال جديد\14-4-2026.xlsx")
print("Sheets found:", res.get("sheets"))
active_s = res.get("sheets")[0]
print(f"Sample rows in {active_s} (total {len(res['sheets_data'][active_s])}):")
for r in res["sheets_data"][active_s][:4]:
    print(" ", [v for v in r if v])
