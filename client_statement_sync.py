import os
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import shutil

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "marble_erp.db")
IS_VERCEL = bool(os.environ.get("VERCEL") or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"))
WORK_DIR = "/tmp/work_dir" if IS_VERCEL else r"I:\خالي أحمد جلال"

font_title = Font(name="Cairo", size=16, bold=True, color="1E3A8A")
font_head = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Cairo", size=11, bold=False, color="0F172A")
font_bold = Font(name="Cairo", size=12, bold=True, color="0F172A")
font_rem = Font(name="Cairo", size=14, bold=True, color="B91C1C")

fill_head = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_green_head = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
fill_tot = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
fill_rem = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)
double_bottom = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='double', color='1E3A8A')
)

align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

def sync_client_statement_excel(client_name: str):
    """
    Synchronizes the client's master Excel statement file on the computer (I:\خالي أحمد جلال)
    and in bundled_files whenever a new invoice or payment is created or modified.
    """
    if not client_name:
        return

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    # Query active invoices (not marked as archived/closed)
    c.execute("""
    SELECT id, invoice_no, invoice_date, net_amount, notes, status 
    FROM invoices 
    WHERE client_name = ? AND (status IS NULL OR status != 'متقفلة (أرشيف قديم)')
    ORDER BY id ASC
    """, (client_name,))
    active_invoices = c.fetchall()

    # Query payments
    c.execute("""
    SELECT id, payment_date, amount_egp, payment_method, notes 
    FROM client_payments 
    WHERE client_name = ?
    ORDER BY id ASC
    """, (client_name,))
    payments = c.fetchall()
    conn.close()

    if client_name == "الحاج صلاح":
        _sync_haj_salah(active_invoices, payments)
    else:
        _sync_generic_client(client_name, active_invoices, payments)

def _sync_haj_salah(active_invoices, payments):
    file_upgraded = os.path.join(WORK_DIR, "الحاج صلاح", "حساب الحاج صلاح - النسخة المطورة 2026.xlsx")
    file_jadeed = os.path.join(WORK_DIR, "الحاج صلاح", "حساب الحاج صلاح جديد 2.xlsx")
    file_bundled = os.path.join(BASE_DIR, "bundled_files", "الحاج صلاح", "حساب الحاج صلاح - النسخة المطورة 2026.xlsx")

    # Load or create workbook
    if os.path.exists(file_upgraded):
        try:
            wb = openpyxl.load_workbook(file_upgraded)
        except Exception:
            wb = openpyxl.Workbook()
    else:
        wb = openpyxl.Workbook()

    sheet_title = "الحاج صلاح جديد"
    if sheet_title in wb.sheetnames:
        del wb[sheet_title]
    ws = wb.create_sheet(sheet_title, 0)
    ws.views.sheetView[0].rightToLeft = True

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = "الحاج صلاح جديد"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    # Invoices Header
    headers_inv = ["م", "بيان الفاتورة", "المبلغ (ج.م)", "التاريخ", "ملاحظات"]
    for col_idx, h in enumerate(headers_inv, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    start_inv_r = 5
    for idx, (inv_id, inv_no, inv_date, amt, note, st) in enumerate(active_invoices, 1):
        curr_r = start_inv_r + idx - 1
        ws.cell(row=curr_r, column=1, value=idx).alignment = align_center
        ws.cell(row=curr_r, column=2, value=inv_no).alignment = align_right
        c_amt = ws.cell(row=curr_r, column=3, value=amt)
        c_amt.alignment = align_center
        c_amt.number_format = '#,##0'
        ws.cell(row=curr_r, column=4, value=str(inv_date)).alignment = align_center
        ws.cell(row=curr_r, column=5, value=note or "").alignment = align_right
        for c in range(1, 6):
            ws.cell(row=curr_r, column=c).font = font_data
            ws.cell(row=curr_r, column=c).border = thin_border

    end_inv_r = start_inv_r + len(active_invoices) - 1 if active_invoices else start_inv_r
    tot_inv_r = end_inv_r + 1

    # Total Invoices Row
    ws.merge_cells(start_row=tot_inv_r, start_column=1, end_row=tot_inv_r, end_column=2)
    ws.cell(row=tot_inv_r, column=1, value="الاجمالي").alignment = align_center
    ws.cell(row=tot_inv_r, column=1).font = font_bold
    ws.cell(row=tot_inv_r, column=1).fill = fill_tot
    ws.cell(row=tot_inv_r, column=2).fill = fill_tot

    c_tot = ws.cell(row=tot_inv_r, column=3, value=f"=SUM(C{start_inv_r}:C{end_inv_r})")
    c_tot.font = font_bold
    c_tot.fill = fill_tot
    c_tot.alignment = align_center
    c_tot.number_format = '#,##0'

    for c in range(1, 6):
        ws.cell(row=tot_inv_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=tot_inv_r, column=c).fill = fill_tot

    # Payments Section
    pay_head_r = tot_inv_r + 2
    headers_pay = ["م", "بيان الدفعة الواصلة", "المبلغ (ج.م)", "التاريخ", "طريقة الدفع"]
    for col_idx, h in enumerate(headers_pay, 1):
        cell = ws.cell(row=pay_head_r, column=col_idx, value=h)
        cell.font = font_head
        cell.fill = fill_green_head
        cell.alignment = align_center
        cell.border = thin_border

    start_pay_r = pay_head_r + 1
    for idx, (p_id, p_date, p_amt, p_meth, p_note) in enumerate(payments, 1):
        curr_r = start_pay_r + idx - 1
        ws.cell(row=curr_r, column=1, value=idx).alignment = align_center
        ws.cell(row=curr_r, column=2, value=p_note or "واصل منه").alignment = align_right
        c_amt = ws.cell(row=curr_r, column=3, value=p_amt)
        c_amt.alignment = align_center
        c_amt.number_format = '#,##0'
        ws.cell(row=curr_r, column=4, value=str(p_date)).alignment = align_center
        ws.cell(row=curr_r, column=5, value=p_meth or "نقداً").alignment = align_center
        for c in range(1, 6):
            ws.cell(row=curr_r, column=c).font = font_data
            ws.cell(row=curr_r, column=c).border = thin_border

    end_pay_r = start_pay_r + len(payments) - 1 if payments else start_pay_r
    tot_pay_r = end_pay_r + 1

    # Total Payments Row
    ws.merge_cells(start_row=tot_pay_r, start_column=1, end_row=tot_pay_r, end_column=2)
    ws.cell(row=tot_pay_r, column=1, value="إجمالي الواصل منه").alignment = align_center
    ws.cell(row=tot_pay_r, column=1).font = font_bold
    ws.cell(row=tot_pay_r, column=1).fill = fill_tot
    ws.cell(row=tot_pay_r, column=2).fill = fill_tot

    c_pay_tot = ws.cell(row=tot_pay_r, column=3, value=f"=SUM(C{start_pay_r}:C{end_pay_r})")
    c_pay_tot.font = font_bold
    c_pay_tot.fill = fill_tot
    c_pay_tot.alignment = align_center
    c_pay_tot.number_format = '#,##0'

    for c in range(1, 6):
        ws.cell(row=tot_pay_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=tot_pay_r, column=c).fill = fill_tot

    # Remaining Balance Row (المتبقي)
    rem_r = tot_pay_r + 2
    ws.merge_cells(start_row=rem_r, start_column=1, end_row=rem_r, end_column=2)
    ws.cell(row=rem_r, column=1, value="المتبقي (الصافي المطلوب)").alignment = align_center
    ws.cell(row=rem_r, column=1).font = font_rem
    ws.cell(row=rem_r, column=1).fill = fill_rem
    ws.cell(row=rem_r, column=2).fill = fill_rem

    c_rem = ws.cell(row=rem_r, column=3, value=f"=C{tot_inv_r}-C{tot_pay_r}")
    c_rem.font = font_rem
    c_rem.fill = fill_rem
    c_rem.alignment = align_center
    c_rem.number_format = '#,##0'

    for c in range(1, 6):
        ws.cell(row=rem_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=rem_r, column=c).fill = fill_rem

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 25

    # Save to disk
    try:
        os.makedirs(os.path.dirname(file_upgraded), exist_ok=True)
        wb.save(file_upgraded)
    except Exception as e:
        print(f"Error saving file_upgraded: {e}")

    try:
        os.makedirs(os.path.dirname(file_bundled), exist_ok=True)
        wb.save(file_bundled)
    except Exception as e:
        print(f"Error saving file_bundled: {e}")

    # Also update sheet 'جديد' in file_jadeed
    if os.path.exists(file_jadeed):
        try:
            wb_j = openpyxl.load_workbook(file_jadeed)
            if "جديد" in wb_j.sheetnames:
                del wb_j["جديد"]
            ws_j = wb_j.create_sheet("جديد", 1)
            ws_j.views.sheetView[0].rightToLeft = True
            ws_j["B1"] = "الحاج صلاح جديد"
            ws_j["B1"].font = font_title
            for idx, (inv_id, inv_no, inv_date, amt, note, st) in enumerate(active_invoices, 4):
                ws_j[f"A{idx}"] = inv_no
                ws_j[f"B{idx}"] = amt
                ws_j[f"C{idx}"] = str(inv_date)
            last_inv_row = 3 + len(active_invoices)
            tot_row = last_inv_row + 2
            ws_j[f"A{tot_row}"] = "الاجمالي"
            ws_j[f"B{tot_row}"] = f"=SUM(B4:B{last_inv_row})"
            pay_start = tot_row + 2
            for idx, (p_id, p_date, p_amt, p_meth, p_note) in enumerate(payments):
                curr = pay_start + idx
                ws_j[f"A{curr}"] = p_note or "واصل منه"
                ws_j[f"B{curr}"] = p_amt
            last_pay_row = pay_start + len(payments) - 1 if payments else pay_start
            rem_row = last_pay_row + 2
            ws_j[f"A{rem_row}"] = "المتبقى"
            ws_j[f"B{rem_row}"] = f"=B{tot_row}-SUM(B{pay_start}:B{last_pay_row})"
            wb_j.save(file_jadeed)
        except Exception as e:
            print(f"Notice on file_jadeed update: {e}")

def _sync_generic_client(client_name, active_invoices, payments):
    file_path = os.path.join(WORK_DIR, client_name, f"حساب {client_name} - النسخة المطورة 2026.xlsx")
    file_bundled = os.path.join(BASE_DIR, "bundled_files", client_name, f"حساب {client_name} - النسخة المطورة 2026.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"كشف حساب {client_name}"
    ws.views.sheetView[0].rightToLeft = True

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = f"كشف حساب وسجل فواتير: {client_name}"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    # Invoices Header
    headers_inv = ["م", "رقم وبيان الفاتورة", "المبلغ (ج.م)", "التاريخ", "ملاحظات"]
    for col_idx, h in enumerate(headers_inv, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_head
        cell.fill = fill_head
        cell.alignment = align_center
        cell.border = thin_border

    start_r = 4
    for idx, (inv_id, inv_no, inv_date, amt, note, st) in enumerate(active_invoices, 1):
        curr_r = start_r + idx - 1
        ws.cell(row=curr_r, column=1, value=idx).alignment = align_center
        ws.cell(row=curr_r, column=2, value=inv_no).alignment = align_right
        c_amt = ws.cell(row=curr_r, column=3, value=amt)
        c_amt.alignment = align_center
        c_amt.number_format = '#,##0.00'
        ws.cell(row=curr_r, column=4, value=str(inv_date)).alignment = align_center
        ws.cell(row=curr_r, column=5, value=note or "").alignment = align_right
        for c in range(1, 6):
            ws.cell(row=curr_r, column=c).font = font_data
            ws.cell(row=curr_r, column=c).border = thin_border

    end_r = start_r + len(active_invoices) - 1 if active_invoices else start_r
    tot_r = end_r + 1

    ws.merge_cells(start_row=tot_r, start_column=1, end_row=tot_r, end_column=2)
    ws.cell(row=tot_r, column=1, value="إجمالي الفواتير").alignment = align_center
    ws.cell(row=tot_r, column=1).font = font_bold
    ws.cell(row=tot_r, column=1).fill = fill_tot
    ws.cell(row=tot_r, column=2).fill = fill_tot

    c_tot = ws.cell(row=tot_r, column=3, value=f"=SUM(C{start_r}:C{end_r})")
    c_tot.font = font_bold
    c_tot.fill = fill_tot
    c_tot.alignment = align_center
    c_tot.number_format = '#,##0.00'

    for c in range(1, 6):
        ws.cell(row=tot_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=tot_r, column=c).fill = fill_tot

    # Payments
    pay_head_r = tot_r + 2
    headers_pay = ["م", "بيان الدفعة الواصلة", "المبلغ (ج.م)", "التاريخ", "طريقة الدفع"]
    for col_idx, h in enumerate(headers_pay, 1):
        cell = ws.cell(row=pay_head_r, column=col_idx, value=h)
        cell.font = font_head
        cell.fill = fill_green_head
        cell.alignment = align_center
        cell.border = thin_border

    start_pay_r = pay_head_r + 1
    for idx, (p_id, p_date, p_amt, p_meth, p_note) in enumerate(payments, 1):
        curr_r = start_pay_r + idx - 1
        ws.cell(row=curr_r, column=1, value=idx).alignment = align_center
        ws.cell(row=curr_r, column=2, value=p_note or "واصل منه").alignment = align_right
        c_amt = ws.cell(row=curr_r, column=3, value=p_amt)
        c_amt.alignment = align_center
        c_amt.number_format = '#,##0.00'
        ws.cell(row=curr_r, column=4, value=str(p_date)).alignment = align_center
        ws.cell(row=curr_r, column=5, value=p_meth or "نقداً").alignment = align_center
        for c in range(1, 6):
            ws.cell(row=curr_r, column=c).font = font_data
            ws.cell(row=curr_r, column=c).border = thin_border

    end_pay_r = start_pay_r + len(payments) - 1 if payments else start_pay_r
    tot_pay_r = end_pay_r + 1

    ws.merge_cells(start_row=tot_pay_r, start_column=1, end_row=tot_pay_r, end_column=2)
    ws.cell(row=tot_pay_r, column=1, value="إجمالي المسدد (واصل منه)").alignment = align_center
    ws.cell(row=tot_pay_r, column=1).font = font_bold
    ws.cell(row=tot_pay_r, column=1).fill = fill_tot
    ws.cell(row=tot_pay_r, column=2).fill = fill_tot

    c_pay_tot = ws.cell(row=tot_pay_r, column=3, value=f"=SUM(C{start_pay_r}:C{end_pay_r})")
    c_pay_tot.font = font_bold
    c_pay_tot.fill = fill_tot
    c_pay_tot.alignment = align_center
    c_pay_tot.number_format = '#,##0.00'

    for c in range(1, 6):
        ws.cell(row=tot_pay_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=tot_pay_r, column=c).fill = fill_tot

    # Remaining
    rem_r = tot_pay_r + 2
    ws.merge_cells(start_row=rem_r, start_column=1, end_row=rem_r, end_column=2)
    ws.cell(row=rem_r, column=1, value="المتبقي (الصافي)").alignment = align_center
    ws.cell(row=rem_r, column=1).font = font_rem
    ws.cell(row=rem_r, column=1).fill = fill_rem
    ws.cell(row=rem_r, column=2).fill = fill_rem

    c_rem = ws.cell(row=rem_r, column=3, value=f"=C{tot_r}-C{tot_pay_r}")
    c_rem.font = font_rem
    c_rem.fill = fill_rem
    c_rem.alignment = align_center
    c_rem.number_format = '#,##0.00'

    for c in range(1, 6):
        ws.cell(row=rem_r, column=c).border = double_bottom
        if c > 3:
            ws.cell(row=rem_r, column=c).fill = fill_rem

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 25

    try:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        wb.save(file_path)
    except Exception as e:
        print(f"Error saving generic client file: {e}")

    try:
        os.makedirs(os.path.dirname(file_bundled), exist_ok=True)
        wb.save(file_bundled)
    except Exception as e:
        print(f"Error saving generic bundled file: {e}")
